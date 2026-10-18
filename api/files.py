"""Uploads: store what a scout attaches and turn it into content blocks the model can read.

    images (png/jpeg/gif/webp)  -> resized JPEG/PNG image block
    PDF                         -> base64 document block (Claude reads text and pages)
    text, markdown, CSV, JSON   -> text document block
    Word (.docx), Excel (.xlsx) -> extracted text document block
    video                       -> the browser extracts frames and uploads them as images

Files are stored per client id under <state>/uploads/<client>/<file_id>/.
"""
import base64
import csv
import io
import json
import re
import uuid
from pathlib import Path

from PIL import Image

MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 150_000
IMAGE_EDGE = 1568  # longest edge Claude reads without downscaling

TEXT_EXT = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm", ".log", ".yaml", ".yml"}
IMAGE_TYPES = {"image/png", "image/jpeg", "image/gif", "image/webp"}


class UploadError(ValueError):
    pass


def _safe_name(name: str) -> str:
    name = Path(name or "file").name
    return re.sub(r"[^\w.\- ()·:]+", "_", name)[:120] or "file"


def _image(data: bytes) -> tuple[bytes, str]:
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception as e:
        raise UploadError("That image couldn't be read. Try PNG or JPEG.") from e
    if max(img.size) > IMAGE_EDGE:
        img.thumbnail((IMAGE_EDGE, IMAGE_EDGE))
    out = io.BytesIO()
    if img.mode in ("RGBA", "LA", "P"):
        img.save(out, "PNG", optimize=True)
        return out.getvalue(), "image/png"
    img.convert("RGB").save(out, "JPEG", quality=85)
    return out.getvalue(), "image/jpeg"


def _docx(data: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets[:5]:
        buf = io.StringIO()
        w = csv.writer(buf)
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 2000:
                break
            w.writerow(["" if v is None else v for v in row])
        out.append(f"## Sheet: {ws.title}\n{buf.getvalue()}")
    return "\n\n".join(out)


def classify(name: str, content_type: str | None, data: bytes) -> tuple[str, bytes, str]:
    """-> (kind, stored bytes, media type). kind is image | pdf | text."""
    if len(data) > MAX_BYTES:
        raise UploadError("Files can be up to 20 MB.")
    ext = Path(name).suffix.lower()
    ctype = (content_type or "").split(";")[0].lower()
    if ctype in IMAGE_TYPES or ext in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
        out, media = _image(data)
        return "image", out, media
    if ctype == "application/pdf" or ext == ".pdf":
        if not data.startswith(b"%PDF"):
            raise UploadError("That file doesn't look like a PDF.")
        return "pdf", data, "application/pdf"
    if ext == ".docx":
        return "text", _docx(data).encode(), "text/plain"
    if ext == ".xlsx":
        return "text", _xlsx(data).encode(), "text/plain"
    if ext in TEXT_EXT or ctype.startswith("text/") or ctype == "application/json":
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        return "text", text.encode(), "text/plain"
    if ctype.startswith("video/"):
        raise UploadError("Videos are sent as still frames; the app extracts them in your browser.")
    raise UploadError(f"Touchline can't read {ext or 'this'} files yet. Try images, PDF, Word, Excel, CSV or text.")


class FileStore:
    """Uploads on local disk under <root>/<client>/<file_id>/ (development)."""

    def __init__(self, root: Path):
        self.root = root

    def _put(self, owner: str, meta: dict, data: bytes) -> None:
        d = self.root / owner / meta["file_id"]
        d.mkdir(parents=True, exist_ok=True)
        (d / "data").write_bytes(data)
        (d / "meta.json").write_text(json.dumps(meta), encoding="utf-8")

    def _load(self, owner: str, file_id: str) -> tuple[dict, bytes] | None:
        d = self.root / owner / file_id
        if not (d / "meta.json").exists():
            return None
        return json.loads((d / "meta.json").read_text(encoding="utf-8")), (d / "data").read_bytes()

    def save(self, owner: str, name: str, content_type: str | None, data: bytes) -> dict:
        kind, stored, media = classify(name, content_type, data)
        meta = {"file_id": uuid.uuid4().hex, "name": _safe_name(name), "kind": kind, "media_type": media,
                "size": len(stored), "note": None}
        if kind == "text" and len(stored.decode("utf-8", errors="replace")) > MAX_TEXT_CHARS:
            meta["note"] = f"Long file: only the first {MAX_TEXT_CHARS:,} characters will be read."
        if kind == "text" and not stored.strip():
            meta["note"] = "No text could be extracted from this file."
        self._put(owner, meta, stored)
        return meta

    def get(self, owner: str, file_id: str) -> tuple[dict, bytes]:
        if not re.fullmatch(r"[0-9a-f]{32}", file_id):
            raise UploadError("unknown file")
        found = self._load(owner, file_id)
        if found is None:
            raise UploadError("unknown file")
        return found

    def blocks(self, owner: str, file_ids: list[str]) -> tuple[list[dict], list[str]]:
        """Content blocks for a message, plus readable labels for the attachment line."""
        blocks, labels = [], []
        for fid in file_ids:
            meta, data = self.get(owner, fid)
            labels.append(f"{meta['name']} ({meta['kind']})")
            if meta["kind"] == "image":
                blocks.append({"type": "image", "source": {"type": "base64", "media_type": meta["media_type"],
                                                           "data": base64.b64encode(data).decode()}})
            elif meta["kind"] == "pdf":
                blocks.append({"type": "document", "title": meta["name"],
                               "source": {"type": "base64", "media_type": "application/pdf",
                                          "data": base64.b64encode(data).decode()}})
            else:
                text = data.decode("utf-8", errors="replace")
                if len(text) > MAX_TEXT_CHARS:
                    text = text[:MAX_TEXT_CHARS] + "\n[truncated]"
                blocks.append({"type": "document", "title": meta["name"],
                               "source": {"type": "text", "media_type": "text/plain", "data": text or "(empty)"}})
        return blocks, labels


class SqlFileStore(FileStore):
    """Uploads in the app database, for hosts without a persistent disk (Render's free tier)."""

    def __init__(self, sql):
        self.sql = sql
        blob = "BYTEA" if sql.dialect == "postgres" else "BLOB"
        sql.script(f"""CREATE TABLE IF NOT EXISTS uploads (
            file_id TEXT PRIMARY KEY, owner TEXT NOT NULL, meta TEXT NOT NULL, data {blob} NOT NULL)""")

    def _put(self, owner: str, meta: dict, data: bytes) -> None:
        self.sql.x("INSERT INTO uploads VALUES (?, ?, ?, ?)", (meta["file_id"], owner, json.dumps(meta), data))

    def _load(self, owner: str, file_id: str) -> tuple[dict, bytes] | None:
        rows = self.sql.q("SELECT meta, data FROM uploads WHERE file_id = ? AND owner = ?", (file_id, owner))
        if not rows:
            return None
        return json.loads(rows[0]["meta"]), bytes(rows[0]["data"])
